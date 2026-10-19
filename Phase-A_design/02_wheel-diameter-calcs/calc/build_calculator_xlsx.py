"""Build wheel_layout_calculator.xlsx (v2, live formulas) for the new chassis ([A-00] Main AMR Assembly).

Mirrors wheel_layout_calc.py: measured CAD positions are inputs; checks and loads are formulas.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "wheel_layout_calculator.xlsx")
F = "Arial"
BLUE, BLACK, GREEN = Font(name=F, color="0000FF"), Font(name=F), Font(name=F, color="008000")
BOLD, TITLE = Font(name=F, bold=True), Font(name=F, bold=True, size=14)
NOTE = Font(name=F, size=9, color="555555")
HEAD, HEAD_FILL = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F4F6F")
SEC_FILL, KEY_FILL, RES_FILL = PatternFill("solid", fgColor="DCE3EA"), PatternFill("solid", fgColor="FFFF00"), PatternFill("solid", fgColor="E8F4EC")
thin = Side(style="thin", color="B7BEC7")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CAD = "[A-00] Main AMR Assembly, read via SolidWorks API 2026-10-20 (GitHub e98996c)"

wb = Workbook()


def header(ws, row, cols):
    for i, h in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = HEAD, HEAD_FILL, Alignment(vertical="center", wrap_text=True), BOX


def section(ws, row, text, n):
    ws.cell(row=row, column=1, value=text).font = BOLD
    for i in range(1, n + 1):
        ws.cell(row=row, column=i).fill = SEC_FILL


# ------------------------------------------------------------------ Inputs
wi = wb.active
wi.title = "Inputs"
wi["A1"] = "Task 2 v2 · Wheel layout calculator · INPUTS"
wi["A1"].font = TITLE
wi["A2"] = ("Frame of [A-00] Main AMR Assembly: origin = robot centre (X = 0, Z = 0), X = width, Y = up, Z = length, mm. "
            "Y = 0 is the underside of the 1 mm Ground part, so the floor surface is Y = 1.0. Edit only BLUE values.")
wi["A2"].font = Font(name=F, italic=True)
header(wi, 4, ["Parameter", "Value", "Unit", "Source / note"])
inputs = [
    ("Measured in the CAD", None),
    ("floor_y", "Floor surface Y (top of Ground part)", 1.0, "mm", CAD),
    ("frame_x", "Frame outer face X (each side)", 262, "mm", CAD + ": 3030 frame 524 wide"),
    ("frame_z", "Frame outer face Z (each end)", 328, "mm", CAD + ": 3030 frame 656 long"),
    ("frame_bottom_y", "Frame bottom Y (underside of bottom profiles)", 46.35, "mm", CAD),
    ("profile_top_y", "Top of bottom profiles = underside of First Plate", 76.35, "mm", CAD),
    ("frame_top_y", "Frame top Y", 166.35, "mm", CAD),
    ("lowest_y", "Lowest chassis part Y (cross joints / tee brackets)", 36.35, "mm", CAD),
    ("axle_y", "Drive axle centre Y (as modelled)", 87.0, "mm", CAD + " ('axe' part)"),
    ("axle_r", "Drive axle radius", 9.7, "mm", CAD + " (19.4 mm shaft)"),
    ("plane_x", "Drive tyre mid-plane X (each side)", 294.3, "mm", CAD + " (tyre X 271.3..317.3)"),
    ("tyre_inner_x", "Drive tyre inner face X", 271.3, "mm", CAD),
    ("tyre_w", "Drive tyre width", 46, "mm", CAD),
    ("tyre_d", "Drive tyre diameter", 172, "mm", CAD + " (rotated tyre box 240.2 / 1.397). Measure the real wheel."),
    ("caster_x", "Caster swivel axis X (each side)", 175, "mm", CAD),
    ("caster_z", "Caster swivel axis Z (each end)", 230.8, "mm", CAD),
    ("caster_d", "Caster wheel diameter", 50.7, "mm", CAD),
    ("caster_lead", "Caster lead (swivel axis to wheel axle)", 23, "mm", CAD),
    ("caster_top_y", "Top of the caster model", 77.85, "mm", CAD),
    ("Design inputs", None),
    ("tyre_defl", "Tyre deflection under load", 1, "mm", "Solid tyre, estimate."),
    ("droop", "Suspension: wheel drop needed below the floor plane", 5, "mm", "Floor dips."),
    ("gap_min", "Minimum axle-to-profile clearance at full drop", 1, "mm", ""),
    ("m_empty", "Robot mass, empty", "='Mass budget'!C16", "kg", "Linked to the Mass budget sheet (CAD has no materials yet)."),
    ("m_payload", "Design payload", 100, "kg", "GP1 sizing basis."),
    ("h_loaded", "CG height above floor, loaded + lift raised", 0.5, "m", "Conservative."),
    ("h_empty", "CG height above floor, empty", 0.15, "m", "Estimate."),
    ("cg_off", "Payload CG offset from centre", 0.05, "m", ""),
    ("accel", "Target acceleration, loaded", 0.5, "m/s2", ""),
    ("slope", "Slope / ramp", 2, "deg", ""),
    ("crr", "Rolling resistance", 0.015, "-", ""),
    ("mu_nom", "Friction, nominal", 0.5, "-", ""),
    ("mu_worst", "Friction, worst", 0.4, "-", ""),
    ("g", "Gravity", 9.81, "m/s2", ""),
    ("preload", "Spring force pressing EACH drive wheel on the floor", 120, "N", "Keep 2 x preload below ~70% of the empty weight."),
    ("floor_tol", "Vertical variation to absorb, +/-", 5, "mm", ""),
    ("force_var", "Allowed +/- force change over that travel", 0.25, "-", ""),
    ("lever_i", "Suspension lever ratio (spring travel / wheel travel)", 1, "-", "1 for a vertical slide."),
    ("caster_sel_h", "Selected caster: total height (datasheet)", 64, "mm", "Example: Revvo 50 mm low-profile, 64 mm, 160 kg (RS 750834)."),
    ("caster_sel_kg", "Selected caster: rating per caster", 160, "kg", ""),
]
names, r = {}, 5
for item in inputs:
    if item[1] is None:
        section(wi, r, item[0], 4)
    else:
        key, label, val, unit, note = item
        wi.cell(row=r, column=1, value=label).font = BLACK
        c = wi.cell(row=r, column=2, value=val)
        c.font = GREEN if isinstance(val, str) and "!" in val else (BLACK if isinstance(val, str) else BLUE)
        wi.cell(row=r, column=3, value=unit).font = BLACK
        wi.cell(row=r, column=4, value=note).font = NOTE
        for col in range(1, 5):
            wi.cell(row=r, column=col).border = BOX
        names[key] = f"Inputs!$B${r}"
    r += 1
for k in ("tyre_d", "m_payload", "preload", "caster_sel_h"):
    wi[names[k].split("!")[1].replace("$", "")].fill = KEY_FILL
for col, w in zip("ABCD", (60, 12, 8, 95)):
    wi.column_dimensions[col].width = w
wi.freeze_panes = "A5"
N = names.__getitem__

# ------------------------------------------------------------------ Mass budget
wm = wb.create_sheet("Mass budget")
wm["A1"] = "Mass budget, empty robot (estimates: the CAD parts have no material assigned yet)"
wm["A1"].font = TITLE
header(wm, 3, ["Item", "Qty", "Mass total (kg)", "Unit mass (kg)", "Source / note"])
mass = [
    ("3030 aluminium profile, total 6.67 m (cut list)", 6.672, 0.9, "~0.9 kg/m typical for 3030; length from 'List of components.xlsx'."),
    ("Brackets, corner joints, cross joints, T-nuts, screws", 1, 1.2, "Estimate."),
    ("First Plate, aluminium 460 x 580 x 3 mm", 1, 2.16, "0.8 m2·mm x 2.7 g/cm3."),
    ("Hoverboard hub motor 6.5 in with tyre", 2, 3.1, "Parallax 6.5 in hub motor listing."),
    ("2 in swivel caster", 4, 0.35, "Estimate."),
    ("Drive suspension module (bracket, slide, spring)", 2, 1.0, "Estimate."),
    ("Lifting mechanism (Task 3)", 1, 5.0, "Estimate; replace with Task 3 mass."),
    ("Battery 36 V ~10 Ah", 1, 3.0, "Estimate."),
    ("Electronics", 1, 2.5, "Estimate."),
    ("Cover, 3 mm sheet (Task 4)", 1, 4.0, "Estimate; Task 4 uses 3 mm sheet."),
]
for i, (item, q, m, note) in enumerate(mass, start=4):
    wm.cell(row=i, column=1, value=item).font = BLACK
    wm.cell(row=i, column=2, value=q).font = BLUE
    wm.cell(row=i, column=4, value=m).font = BLUE
    wm.cell(row=i, column=3, value=f"=B{i}*D{i}").font = BLACK
    wm.cell(row=i, column=5, value=note).font = NOTE
    for col in range(1, 6):
        wm.cell(row=i, column=col).border = BOX
    wm[f"C{i}"].number_format = "0.00"
wm["A14"], wm["C14"] = "Subtotal", "=SUM(C4:C13)"
wm["A15"], wm["C15"] = "Design margin (10%)", "=0.1*C14"
wm["A16"], wm["C16"] = "Empty robot mass used in calculations", "=C14+C15"
for rr in (14, 15, 16):
    wm[f"A{rr}"].font, wm[f"C{rr}"].font, wm[f"C{rr}"].number_format = BOLD, BOLD, "0.0"
wm["C16"].fill = RES_FILL
for col, w in zip("ABCDE", (55, 8, 16, 15, 70)):
    wm.column_dimensions[col].width = w

# ------------------------------------------------------------------ Results
wr = wb.create_sheet("Results")
wr["A1"] = "Task 2 v2 · RESULTS (formulas, do not edit)"
wr["A1"].font = TITLE
header(wr, 3, ["Result", "Value", "Unit", "Meaning / check"])
res, rr = {}, 4


def put(key, label, formula, unit, note, fmt="0.00"):
    global rr
    wr.cell(row=rr, column=1, value=label).font = BLACK
    c = wr.cell(row=rr, column=2, value=formula)
    c.font, c.number_format, c.fill = BLACK, fmt, RES_FILL
    wr.cell(row=rr, column=3, value=unit).font = BLACK
    wr.cell(row=rr, column=4, value=note).font = NOTE
    for col in range(1, 5):
        wr.cell(row=rr, column=col).border = BOX
    res[key] = f"$B${rr}"
    rr += 1


def sec(t):
    global rr
    section(wr, rr, t, 4)
    rr += 1


R_ = res.__getitem__
sec("1. Vertical alignment (as modelled)")
put("contact", "Drive tyre bottom Y (axle Y - tyre D/2)", f"={N('axle_y')}-{N('tyre_d')}/2", "mm", "Should equal the floor Y")
put("contact_err", "Drive tyre vs floor", f"={R_('contact')}-{N('floor_y')}", "mm", "0 = drive wheels and casters on the same plane")
put("caster_mount_h", "Floor to caster mounting face (underside of First Plate)", f"={N('profile_top_y')}-{N('floor_y')}", "mm", "")
put("caster_overlap", "Caster model top above the mounting face", f"={N('caster_top_y')}-{N('profile_top_y')}", "mm", ">0 = the caster model pokes into the plate")
put("clear_frame", "Ground clearance under the frame", f"={N('frame_bottom_y')}-{N('floor_y')}", "mm", "")
put("clear_lowest", "Ground clearance under the lowest chassis part", f"={N('lowest_y')}-{N('floor_y')}", "mm", "")
put("tyre_above_top", "Tyre top above the frame top", f"={N('axle_y')}+{N('tyre_d')}/2-{N('frame_top_y')}", "mm", ">0 = needs wheel wells in the cover")
put("tyre_gap", "Gap tyre inner face to frame outer face", f"={N('tyre_inner_x')}-{N('frame_x')}", "mm", "")
sec("2. Axle clearance over the bottom side profile (the axle passes over it at Z = 0)")
put("gap_model", "Gap as modelled (axle unloaded)", f"={N('axle_y')}-{N('axle_r')}-{N('profile_top_y')}", "mm", "")
put("gap_loaded", "Gap under load (tyre compressed)", f"={N('floor_y')}+{N('tyre_d')}/2-{N('tyre_defl')}-{N('axle_r')}-{N('profile_top_y')}", "mm", "<0 = the axle sits on the profile")
put("gap_need", "Gap needed (droop + clearance)", f"={N('droop')}+{N('gap_min')}", "mm", "")
put("lower", "Lower the frame by (whole mm)", f"=MAX(0,ROUNDUP({R_('gap_need')}-{R_('gap_loaded')},0))", "mm", "By shortening the installed casters", "0")
put("new_caster_h", "New floor-to-caster-face height", f"={R_('caster_mount_h')}-{R_('lower')}", "mm", "Caster + spacer must equal this")
put("spacer", "Spacer for the selected caster", f"={R_('new_caster_h')}-{N('caster_sel_h')}", "mm", "")
put("spacer_ok", "Selected caster fits?", f'=IF({R_("spacer")}<0,"NO: too tall","OK")', "", "", "@")
put("new_clear", "New clearance under the lowest part", f"={R_('clear_lowest')}-{R_('lower')}", "mm", "")
put("new_tyre_above", "New tyre top above the frame top", f"={R_('tyre_above_top')}+{R_('lower')}", "mm", "")
sec("3. Plan geometry")
put("track", "Track (tyre mid-planes)", f"=2*{N('plane_x')}", "mm", "")
put("width", "Width over the tyres", f"=2*({N('tyre_inner_x')}+{N('tyre_w')})", "mm", "")
put("c_track", "Caster track (X)", f"=2*{N('caster_x')}", "mm", "")
put("c_base", "Caster wheelbase (Z)", f"=2*{N('caster_z')}", "mm", "")
put("swivel", "Caster swivel envelope radius", f"={N('caster_lead')}+{N('caster_d')}/2", "mm", "")
put("rot", "In-place rotation radius (frame corners)", f"=SQRT({N('frame_x')}^2+{N('frame_z')}^2)", "mm", "Add the cover thickness")
sec("4. Stability (casters swivelled inwards)")
put("ax", "Caster inner contact X", f"={N('caster_x')}-{N('caster_lead')}", "mm", "")
put("az", "Caster inner contact Z", f"={N('caster_z')}-{N('caster_lead')}", "mm", "")
put("m_side", "Margin to the caster-to-drive-wheel edge",
    f"=ABS({N('plane_x')}*{R_('az')})/SQRT(({N('plane_x')}-{R_('ax')})^2+{R_('az')}^2)/1000", "m", "", "0.000")
put("m_long", "Margin to the end edge", f"={R_('az')}/1000", "m", "", "0.000")
put("m_min", "Smallest margin", f"=MIN({R_('m_side')},{R_('m_long')})", "m", "", "0.000")
put("tip_l", "Tip-over acceleration limit, loaded", f"={N('g')}*({R_('m_min')}-{N('cg_off')})/{N('h_loaded')}", "m/s2", "")
put("tip_e", "Tip-over acceleration limit, empty", f"={N('g')}*{R_('m_min')}/{N('h_empty')}", "m/s2", "")
sec("5. Drive, traction and casters")
put("m_l", "Loaded mass", f"={N('m_empty')}+{N('m_payload')}", "kg", "", "0.0")
put("r_m", "Loaded wheel radius", f"=({N('tyre_d')}/2-{N('tyre_defl')})/1000", "m", "", "0.0000")
put("f_slope", "Drive force, loaded, slope, target acceleration", f"={R_('m_l')}*({N('g')}*{N('crr')}*COS(RADIANS({N('slope')}))+{N('g')}*SIN(RADIANS({N('slope')}))+{N('accel')})", "N", "")
put("t_slope", "Torque per wheel, slope case", f"={R_('f_slope')}/2*{R_('r_m')}", "N·m", "Hub-motor requirement")
put("rpm", "Wheel speed at 1 m/s", f"=60/(PI()*{N('tyre_d')}/1000)", "rpm", "", "0.0")
put("pre_frac", "Drive preload / empty weight", f"=2*{N('preload')}/({N('m_empty')}*{N('g')})", "-", "Keep below ~70%", "0%")
put("c_empty", "Load per caster, empty", f"=({N('m_empty')}*{N('g')}-2*{N('preload')})/4", "N", "Must stay > ~15 N", "0.0")
put("c_static", "Load per caster, loaded, static", f"=({R_('m_l')}*{N('g')}-2*{N('preload')})/4", "N", "", "0.0")
put("c_worst", "Load per caster, loaded, worst", f"={R_('c_static')}+{R_('m_l')}*{N('accel')}*{N('h_loaded')}/({R_('c_base')}/1000)/2+{R_('m_l')}*{N('g')}*{N('cg_off')}/({R_('c_base')}/1000)/2", "N", "", "0.0")
put("c_req", "Required caster rating (x1.5)", f"={R_('c_worst')}*1.5/{N('g')}", "kg", "", "0.0")
put("c_ok", "Selected caster strong enough?", f'=IF({N("caster_sel_kg")}>={R_("c_req")},"OK","NO")', "", "", "@")
put("a_nom", "Max acceleration before slip, loaded, flat (mu nominal)", f"=(2*{N('mu_nom')}*{N('preload')}-{R_('m_l')}*{N('g')}*{N('crr')})/{R_('m_l')}", "m/s2", "")
put("a_worst", "Max acceleration before slip, loaded, flat (mu worst)", f"=(2*{N('mu_worst')}*{N('preload')}-{R_('m_l')}*{N('g')}*{N('crr')})/{R_('m_l')}", "m/s2", "Motor-controller limit")
put("a_slope", "Max acceleration before slip, loaded, slope (mu worst)", f"=(2*{N('mu_worst')}*{N('preload')}-{R_('m_l')}*{N('g')}*({N('crr')}*COS(RADIANS({N('slope')}))+SIN(RADIANS({N('slope')}))))/{R_('m_l')}", "m/s2", "")
put("k", "Max suspension rate at the wheel", f"={N('force_var')}*{N('preload')}/{N('floor_tol')}", "N/mm", "")
put("comp", "Preload compression at the wheel", f"={N('preload')}/{R_('k')}", "mm", "")
put("k_spring", "Spring rate at the spring", f"={R_('k')}/{N('lever_i')}^2", "N/mm", "")
for col, w in zip("ABCD", (62, 14, 8, 55)):
    wr.column_dimensions[col].width = w
wr.freeze_panes = "A4"

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold if c.font else False, color=c.font.color if c.font else None,
                              italic=c.font.italic if c.font else False, size=c.font.size if c.font else 11)
wb.save(OUT)
print("saved", OUT)
